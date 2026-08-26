"""End-to-end verification: Knowledge Profile rebuild + document profile-context.

Executes docs/architecture/knowledge-profile-visibility-and-async-e2e-verification-plan.md
§6.2 against a live stack: docker compose services (postgres/redis/minio), a
running API server and an RQ worker, both started with the fake model
providers:

    M2_EMBEDDING_PROVIDER=fake M3_LLM_PROVIDER=fake \
    M4_OCR_PROVIDER=fake M4_VISION_PROVIDER=fake

Phases (bash stops/starts the worker between phases for deterministic control):

    phase1  upload/build/publish + rebuild branches a/b + profile-context + auth
    phase2  [worker STOPPED] branch (c): merge intent into a queued build
    phase3  [worker running] merged build completes -> chained PROFILE_REFRESH
    phase4  [worker STOPPED] queue a refresh, then cancel the release (SQL)
    phase5  [worker running] task FAILED, then retry completes the rebuild
    phase6  knowledge base deletion flow

State (knowledge/task ids) is carried between phases via --state-file.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import sys
import tempfile
import time
from pathlib import Path
from typing import Any
from uuid import uuid4

import asyncpg
import httpx
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

API = "http://127.0.0.1:8765"
TOKEN = "change-me-for-local-development"
DB_DSN = "postgresql://rag:rag-dev-password@127.0.0.1:15432/rag_platform"
STATE_FILE = Path("/tmp/e2e_profile_state.json")

EVIDENCE: list[str] = []


def log(message: str) -> None:
    print(message, flush=True)
    EVIDENCE.append(message)


def check(condition: bool, label: str) -> None:
    status = "PASS" if condition else "FAIL"
    log(f"[{status}] {label}")
    if not condition:
        raise AssertionError(f"verification failed: {label}")


class Api:
    def __init__(self, client: httpx.Client) -> None:
        self.client = client

    def request(self, method: str, path: str, **kwargs: Any) -> httpx.Response:
        headers = kwargs.pop("headers", {})
        headers.setdefault("Authorization", f"Bearer {TOKEN}")
        return self.client.request(method, f"{API}{path}", headers=headers, **kwargs)


def wait_task(api: Api, task_id: str, timeout: float = 120.0) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = api.request("GET", f"/api/v1/tasks/{task_id}")
        if response.status_code != 200:
            raise AssertionError(f"task poll failed: {response.status_code} {response.text}")
        task = response.json()
        if task["status"] in {"SUCCEEDED", "FAILED"}:
            return task
        time.sleep(0.4)
    raise AssertionError(f"task {task_id} did not reach a terminal state in {timeout}s")


def make_docx(path: Path, heading: str, paragraph: str) -> None:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(paragraph)
    document.save(str(path))


def load_state() -> dict[str, Any]:
    if STATE_FILE.exists():
        return json.loads(STATE_FILE.read_text())
    return {}


def save_state(state: dict[str, Any]) -> None:
    STATE_FILE.write_text(json.dumps(state, indent=2, ensure_ascii=False))


# ---------------------------------------------------------------------------


def phase1(api: Api, state: dict[str, Any]) -> None:
    log("=== Phase 1: upload -> candidate build -> profile generation -> branches ===")

    # §6.2.1 — create KB + upload two DOCX; HTTP must return fast (202) while
    # the heavy work happens in the worker.
    kb_name = f"e2e-profile-{uuid4().hex[:8]}"
    created = api.request("POST", "/api/v1/knowledge-bases", json={"name": kb_name, "description": None})
    check(created.status_code == 201, f"create knowledge base -> 201 ({kb_name})")
    kb_id = created.json()["id"]
    state["kb_id"] = kb_id
    state["kb_name"] = kb_name
    log(f"  knowledge_id={kb_id}")

    with tempfile.TemporaryDirectory() as tmp:
        doc_paths = [
            Path(tmp) / "first.docx",
            Path(tmp) / "second.docx",
        ]
        make_docx(doc_paths[0], "PDC 升级步骤", "PDC 升级前请断开主电源，并记录当前版本号。")
        make_docx(doc_paths[1], "风扇更换流程", "更换风扇前请关闭设备电源并等待五分钟。")
        document_ids: list[str] = []
        for path in doc_paths:
            started = time.perf_counter()
            upload = api.request(
                "POST",
                f"/api/v1/knowledge-bases/{kb_id}/documents",
                headers={"Idempotency-Key": f"e2e-{uuid4()}"},
                files={"file": (path.name, path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"display_name": path.stem, "request_build": "true"},
            )
            elapsed = time.perf_counter() - started
            check(upload.status_code == 202, f"upload {path.name} -> 202 (HTTP took {elapsed:.2f}s, heavy work deferred to worker)")
            document_ids.append(upload.json()["document"]["id"])
    state["document_ids"] = document_ids
    build = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/indexes/build",
                        headers={"Idempotency-Key": f"e2e-{uuid4()}"},
                        json={"reason": "MANUAL", "activate_on_success": False}).json()
    if build["coalesced"]:
        # The uploads already started a candidate build; join that task.
        build_tasks = api.request("GET", f"/api/v1/indexes/{build['index_id']}/tasks").json()
        build["task_id"] = next(
            t["id"] for t in build_tasks
            if t.get("task_type", "INDEX_BUILD") == "INDEX_BUILD"
            and t["status"] in {"QUEUED", "RUNNING"}
        )
        check(True, f"manual build coalesced into the upload-triggered candidate (index {build['index_id']}, task {build['task_id']})")
    else:
        check(build["requested"] is True, f"index build submitted (index {build['index_id']})")

    # §6.2.2 — wait for the candidate chain to settle: a coalesced upload
    # creates a followup build, so wait until no index is BUILDING, then check
    # the latest candidate's INDEX_BUILD and PROFILE_GENERATE tasks.
    def wait_stable_candidate(timeout: float = 180.0) -> dict[str, Any]:
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            indexes = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}/indexes").json()
            if not any(i["status"] == "BUILDING" for i in indexes):
                latest = max(indexes, key=lambda i: i["version"])
                tasks = api.request("GET", f"/api/v1/indexes/{latest['id']}/tasks").json()
                profile = next(
                    (t for t in tasks if str(t.get("task_type", "")).startswith("PROFILE_")), None
                )
                if profile is not None:
                    return latest
            time.sleep(0.5)
        raise AssertionError("candidate build chain did not settle with a profile task")

    latest = wait_stable_candidate()
    index_task = wait_task(api, build["task_id"])
    check(index_task["status"] == "SUCCEEDED",
          f"INDEX_BUILD SUCCEEDED (stage={index_task['stage']}, progress={index_task['progress']}, attempt={index_task['attempt']})")
    log(f"  index task metadata: {json.dumps(index_task['metadata'], ensure_ascii=False)}")
    log(f"  stable candidate index: v{latest['version']} ({latest['id']})")
    index_tasks = api.request("GET", f"/api/v1/indexes/{latest['id']}/tasks").json()
    profile_submission = next(
        (t for t in index_tasks if str(t.get("task_type", "")).startswith("PROFILE_")), None
    )
    assert profile_submission is not None, "no PROFILE task submitted for the candidate build"
    profile_task = wait_task(api, profile_submission["id"])
    check(profile_task["task_type"] == "PROFILE_GENERATE", "build-time profile task is PROFILE_GENERATE")
    check(profile_task["status"] == "SUCCEEDED",
          f"PROFILE_GENERATE SUCCEEDED (generated_count={profile_task['metadata'].get('generated_count')})")

    releases = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}/releases").json()
    open_release = next(r for r in releases if r["status"] not in {"PUBLISHED", "CANCELLED"})
    detail = api.request("GET", f"/api/v1/releases/{open_release['id']}").json()
    check(len(detail["drafts"]) == 2, f"release {open_release['id']} exposes {len(detail['drafts'])} reviewable drafts")
    state["release_id"] = open_release["id"]
    state["index_id"] = open_release["index_id"]

    # §6.2.3 branch (a) — READY editable candidate -> PROFILE_REFRESH.
    first = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                        headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(first["task_type"] == "PROFILE_REFRESH" and first["replayed"] is False,
          f"branch READY-candidate -> new PROFILE_REFRESH task {first['task_id']}")
    second = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                         headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(second["task_id"] == first["task_id"] and second["replayed"] is True,
          "repeated click returns the SAME active task (replayed), no parallel generation")
    refresh = wait_task(api, first["task_id"])
    check(refresh["status"] == "SUCCEEDED", "PROFILE_REFRESH SUCCEEDED (deduped double click)")
    log(f"  refresh metadata: {json.dumps(refresh['metadata'], ensure_ascii=False)}")

    # The refresh replaced the automatic drafts; re-read them before review.
    detail = api.request("GET", f"/api/v1/releases/{open_release['id']}").json()

    # §3.2 / §4.2 — document profile-context before publication.
    doc_id = document_ids[0]
    context = api.request("GET", f"/api/v1/documents/{doc_id}/profile-context").json()
    check(context["published_profile"] is None, "profile-context: no published profile before first publish")
    check(context["candidate_draft"] is not None and context["candidate_draft"]["status"] in {"DRAFT", "APPROVED"},
          f"profile-context: candidate draft visible (operation={context['candidate_draft']['operation']}, quality={context['candidate_draft']['quality_status']})")
    check("reasoning" not in json.dumps(context), "profile-context does not leak LLM reasoning or storage paths")

    # §6.4 — auth boundaries for the new read endpoint.
    no_token = api.client.get(f"{API}/api/v1/documents/{doc_id}/profile-context")
    check(no_token.status_code == 401, "profile-context without token -> 401")
    bad_token = api.client.get(f"{API}/api/v1/documents/{doc_id}/profile-context",
                               headers={"Authorization": "Bearer wrong"})
    check(bad_token.status_code == 401, "profile-context with wrong token -> 401")
    missing = api.request("GET", f"/api/v1/documents/{uuid4()}/profile-context")
    check(missing.status_code == 404 and missing.json()["error"]["code"] == "DOCUMENT_NOT_FOUND",
          "profile-context for unknown document -> DOCUMENT_NOT_FOUND")

    # Publish flow: approve drafts -> validate -> publish atomically.
    for draft in detail["drafts"]:
        if draft["quality_status"] in {"REVIEW_REQUIRED", "FALLBACK"}:
            patch = api.request("PATCH", f"/api/v1/profile-drafts/{draft['id']}", json={"problem_code": draft["problem_code"]})
            check(patch.status_code == 200, f"revised REVIEW_REQUIRED draft {draft['id']} -> MANUAL_REVIEWED")
        approved = api.request("PATCH", f"/api/v1/profile-drafts/{draft['id']}", json={"status": "APPROVED"})
        check(approved.status_code == 200, f"draft {draft['id']} approved")
    validation = api.request("POST", f"/api/v1/releases/{open_release['id']}/validate").json()
    check(validation["publishable"] is True, f"release validation publishable (review_required={validation.get('review_required_count')})")
    published = api.request("POST", f"/api/v1/releases/{open_release['id']}/publish")
    check(published.status_code == 200, "release published atomically (index + profiles)")
    state["published"] = True

    context_after = api.request("GET", f"/api/v1/documents/{doc_id}/profile-context").json()
    check(context_after["published_profile"] is not None,
          f"profile-context after publish: live profile problem_code={context_after['published_profile']['problem_code']}")
    check(context_after["candidate_draft"] is None, "profile-context after publish: candidate draft no longer reported as unpublished")

    # §6.2.3 branch (b) — only an Active Index, no editable candidate ->
    # orchestrated build carrying the refresh intent.
    branch_b = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                           headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(branch_b["task_type"] == "INDEX_BUILD" and branch_b["replayed"] is False,
          f"branch only-active-index -> orchestrated INDEX_BUILD {branch_b['task_id']}")
    build_task = wait_task(api, branch_b["task_id"])
    check(build_task["status"] == "SUCCEEDED", "orchestrated build SUCCEEDED")
    chained_id = build_task["metadata"].get("profile_task_id")
    check(isinstance(chained_id, str) and chained_id, f"build task chains profile task {chained_id} in metadata")
    chained = wait_task(api, chained_id)
    check(chained["task_type"] == "PROFILE_REFRESH" and chained["status"] == "SUCCEEDED",
          "chained task is PROFILE_REFRESH and SUCCEEDED (build-then-refresh orchestration)")

    save_state(state)
    log("=== Phase 1 complete ===")


def phase2(api: Api, state: dict[str, Any]) -> None:
    """Worker is STOPPED for this phase (bash orchestrates)."""
    log("=== Phase 2: branch BUILDING — merge refresh intent into a queued build (worker stopped) ===")
    kb_id = state["kb_id"]
    build = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/indexes/build",
                        headers={"Idempotency-Key": f"e2e-{uuid4()}"},
                        json={"reason": "MANUAL", "activate_on_success": False}).json()
    if build["coalesced"]:
        raise AssertionError(f"expected a fresh build, got coalesced: {build}")
    state["branch_c_task_id"] = build["task_id"]
    state["branch_c_index_id"] = build["index_id"]

    # §6.2.3 branch (c) — candidate index is BUILDING: the rebuild must merge
    # into that build's task instead of creating a parallel generation.
    merged = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                         headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(merged["task_id"] == build["task_id"] and merged["task_type"] == "INDEX_BUILD"
          and merged["replayed"] is True,
          f"branch BUILDING -> SAME INDEX_BUILD task {build['task_id']} returned (replayed), no parallel task")
    task = api.request("GET", f"/api/v1/tasks/{build['task_id']}").json()
    check(task["metadata"].get("profile_refresh_after_build") is True,
          "merged intent persisted in task_record.metadata (profile_refresh_after_build)")
    save_state(state)
    log("=== Phase 2 complete (restart the worker, then run phase3) ===")


def phase3(api: Api, state: dict[str, Any]) -> None:
    log("=== Phase 3: merged intent converts the post-build task to PROFILE_REFRESH ===")
    build_task = wait_task(api, state["branch_c_task_id"])
    check(build_task["status"] == "SUCCEEDED",
          f"merged-intent build SUCCEEDED (stage={build_task['stage']}, attempt={build_task['attempt']})")
    chained_id = build_task["metadata"].get("profile_task_id")
    check(isinstance(chained_id, str) and chained_id, f"build task chains profile task {chained_id} in metadata")
    chained = wait_task(api, chained_id)
    check(chained["task_type"] == "PROFILE_REFRESH" and chained["status"] == "SUCCEEDED",
          "post-build task is PROFILE_REFRESH (not PROFILE_GENERATE) and SUCCEEDED — no duplicate generation")
    save_state(state)
    log("=== Phase 3 complete (stop the worker, then run phase4) ===")


def phase4(api: Api, state: dict[str, Any]) -> None:
    """Worker is STOPPED for this phase (bash orchestrates)."""
    log("=== Phase 4: queue a refresh, then invalidate its candidate (worker stopped) ===")
    kb_id = state["kb_id"]
    submission = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                             headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(submission["task_type"] == "PROFILE_REFRESH", f"queued PROFILE_REFRESH {submission['task_id']} while worker is stopped")
    task = api.request("GET", f"/api/v1/tasks/{submission['task_id']}").json()
    check(task["status"] == "QUEUED", "task remains QUEUED (no worker running)")

    # Cancel the candidate the refresh targets (simulates a superseded
    # candidate) so the worker's execution must fail loudly.
    async def cancel() -> str:
        conn = await asyncpg.connect(DB_DSN)
        try:
            row = await conn.fetchrow(
                "SELECT r.id FROM knowledge_release r "
                "JOIN knowledge_index i ON i.id = r.index_id "
                "WHERE r.kb_id=$1 AND r.status NOT IN ('PUBLISHED','CANCELLED') "
                "ORDER BY i.version DESC LIMIT 1", kb_id)
            await conn.execute("UPDATE knowledge_release SET status='CANCELLED' WHERE id=$1", row["id"])
            return str(row["id"])
        finally:
            await conn.close()

    cancelled = asyncio.run(cancel())
    log(f"  cancelled open release {cancelled} behind the queued refresh")
    state["failing_task_id"] = submission["task_id"]
    save_state(state)
    log("=== Phase 4 complete (restart the worker, then run phase5) ===")


def phase5(api: Api, state: dict[str, Any]) -> None:
    log("=== Phase 5: failure surfacing + retry completes the rebuild ===")
    failed = wait_task(api, state["failing_task_id"], timeout=60)
    check(failed["status"] == "FAILED",
          f"invalidated refresh FAILED with readable error (error_code={failed['error_code']}, message={failed['error_message']!r})")

    # Retry: a fresh rebuild submission must succeed end-to-end.
    retried = api.request("POST", f"/api/v1/knowledge-bases/{state['kb_id']}/profiles/rebuild",
                          headers={"Idempotency-Key": f"e2e-{uuid4()}"}).json()
    check(retried["task_type"] == "INDEX_BUILD",
          f"retry rebuild -> new orchestrated build {retried['task_id']} (candidate was cancelled)")
    build_task = wait_task(api, retried["task_id"])
    check(build_task["status"] == "SUCCEEDED", "retry build SUCCEEDED")
    chained = wait_task(api, build_task["metadata"]["profile_task_id"])
    check(chained["status"] == "SUCCEEDED",
          f"retry chained PROFILE_REFRESH SUCCEEDED (generated_count={chained['metadata'].get('generated_count')}) — retry completes")
    save_state(state)
    log("=== Phase 5 complete ===")


def phase6(api: Api, state: dict[str, Any]) -> None:
    log("=== Phase 6: knowledge base deletion with durable storage cleanup ===")
    kb_id = state["kb_id"]

    def list_minio_objects() -> list[str]:
        from minio import Minio

        client = Minio("127.0.0.1:19000", access_key="rag-minio", secret_key="rag-minio-password", secure=False)
        return [
            item.object_name
            for item in client.list_objects("rag-documents", prefix=f"knowledge-bases/{kb_id}/", recursive=True)
        ]

    objects_before = asyncio.run(asyncio.to_thread(list_minio_objects))
    check(len(objects_before) > 0, f"{len(objects_before)} MinIO source objects present before deletion")

    first = api.request("DELETE", f"/api/v1/knowledge-bases/{kb_id}").json()
    check(first["task_type"] == "KNOWLEDGE_BASE_DELETE" and first["replayed"] is False,
          f"deletion task submitted ({first['task_id']})")
    kb = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}").json()
    check(kb["status"] == "DELETING", "knowledge base visible as DELETING while cleanup runs")

    repeat = api.request("DELETE", f"/api/v1/knowledge-bases/{kb_id}").json()
    check(repeat["task_id"] == first["task_id"] and repeat["replayed"] is True,
          "repeated delete click replays the SAME deletion task")

    done = wait_task(api, first["task_id"], timeout=120)
    check(done["status"] == "SUCCEEDED",
          f"deletion SUCCEEDED (cleanup_completed={done['metadata'].get('cleanup_completed')}, total={done['metadata'].get('cleanup_total')})")

    missing = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}")
    check(missing.status_code == 404, "knowledge base gone after deletion")

    # task_record survives the KB purge (ON DELETE SET NULL + metadata snapshot)
    still = api.request("GET", f"/api/v1/tasks/{first['task_id']}").json()
    check(still["status"] == "SUCCEEDED" and still["metadata"].get("knowledge_id") == kb_id,
          "GET /tasks/{id} still returns SUCCEEDED after physical deletion (kb_id SET NULL, snapshot kept)")

    objects_after = asyncio.run(asyncio.to_thread(list_minio_objects))
    check(objects_after == [], "all MinIO objects under the KB prefix removed")

    # §6.5 — no unexplained leftover records for this KB.
    async def leftovers() -> int:
        conn = await asyncpg.connect(DB_DSN)
        try:
            counts = await conn.fetchrow(
                "SELECT (SELECT count(*) FROM document_source WHERE kb_id=$1) AS documents, "
                "(SELECT count(*) FROM knowledge_index WHERE kb_id=$1) AS indexes, "
                "(SELECT count(*) FROM task_record t JOIN knowledge_index i ON i.id=t.index_id "
                " WHERE i.kb_id=$1 AND t.status IN ('QUEUED','RUNNING')) AS active_tasks",
                kb_id,
            )
            return sum(counts.values())
        finally:
            await conn.close()

    check(asyncio.run(leftovers()) == 0, "no leftover documents/indexes/active tasks for the deleted KB")
    log("=== Phase 6 complete ===")


def phase_browser_setup(api: Api, state: dict[str, Any]) -> None:
    """Prepare stable data for the §6.3 browser acceptance walkthrough."""
    log("=== Browser setup: published + candidate + no-profile states ===")
    kb_name = f"浏览器验收-{uuid4().hex[:6]}"
    created = api.request("POST", "/api/v1/knowledge-bases", json={"name": kb_name, "description": "浏览器验收数据"})
    kb_id = created.json()["id"]
    state["kb_id"] = kb_id
    state["kb_name"] = kb_name
    log(f"  knowledge base: {kb_name} ({kb_id})")

    with tempfile.TemporaryDirectory() as tmp:
        uploads = [
            ("PDC 升级步骤.docx", "PDC 升级前请断开主电源，并记录当前版本号。"),
            ("风扇更换流程.docx", "更换风扇前请关闭设备电源并等待五分钟。"),
        ]
        for index, (name, body) in enumerate(uploads):
            path = Path(tmp) / name
            make_docx(path, name.removesuffix(".docx"), body)
            upload = api.request(
                "POST",
                f"/api/v1/knowledge-bases/{kb_id}/documents",
                headers={"Idempotency-Key": f"browser-{uuid4()}"},
                files={"file": (name, path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"display_name": name.removesuffix(".docx"), "request_build": "true" if index == 0 else "true"},
            )
            check(upload.status_code == 202, f"uploaded {name}")

    # Wait for the candidate chain to settle, then approve + publish.
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        indexes = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}/indexes").json()
        if not any(i["status"] == "BUILDING" for i in indexes):
            latest = max(indexes, key=lambda i: i["version"])
            tasks = api.request("GET", f"/api/v1/indexes/{latest['id']}/tasks").json()
            profile = next((t for t in tasks if str(t.get("task_type", "")).startswith("PROFILE_")), None)
            if profile is not None and profile["status"] in {"SUCCEEDED", "FAILED"}:
                break
        time.sleep(0.5)
    releases = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}/releases").json()
    release = next(r for r in releases if r["status"] not in {"PUBLISHED", "CANCELLED"})
    detail = api.request("GET", f"/api/v1/releases/{release['id']}").json()
    for draft in detail["drafts"]:
        api.request("PATCH", f"/api/v1/profile-drafts/{draft['id']}", json={"problem_code": draft["problem_code"]})
        approved = api.request("PATCH", f"/api/v1/profile-drafts/{draft['id']}", json={"status": "APPROVED"})
        check(approved.status_code == 200, f"draft approved for {draft.get('document_name')}")
    api.request("POST", f"/api/v1/releases/{release['id']}/validate")
    published = api.request("POST", f"/api/v1/releases/{release['id']}/publish")
    check(published.status_code == 200, "published (documents 1+2 now have live profiles)")

    # Rebuild -> new candidate with drafts (published + candidate visible).
    rebuild = api.request("POST", f"/api/v1/knowledge-bases/{kb_id}/profiles/rebuild",
                          headers={"Idempotency-Key": f"browser-{uuid4()}"}).json()
    build_task = wait_task(api, rebuild["task_id"])
    check(build_task["status"] == "SUCCEEDED", "candidate rebuild build finished")
    chained = wait_task(api, build_task["metadata"]["profile_task_id"])
    check(chained["status"] == "SUCCEEDED", "candidate drafts regenerated")

    # A document with neither published profile nor draft.
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "未构建文档.docx"
        make_docx(path, "未构建文档", "此文档未参与构建，用于验证无 Profile 状态。")
        upload = api.request(
            "POST",
            f"/api/v1/knowledge-bases/{kb_id}/documents",
            headers={"Idempotency-Key": f"browser-{uuid4()}"},
            files={"file": ("未构建文档.docx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"display_name": "未构建文档", "request_build": "false"},
        )
        check(upload.status_code == 202, "uploaded 未构建文档 (request_build=false, no profile expected)")

    # A second, disabled knowledge base for the disabled-button check.
    disabled = api.request("POST", "/api/v1/knowledge-bases", json={"name": f"已停用库-{uuid4().hex[:6]}", "description": None}).json()
    state["disabled_kb_id"] = disabled["id"]

    async def disable_kb() -> None:
        conn = await asyncpg.connect(DB_DSN)
        try:
            await conn.execute("UPDATE knowledge_base SET status='DISABLED' WHERE id=$1", disabled["id"])
        finally:
            await conn.close()

    asyncio.run(disable_kb())
    documents = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}/documents").json()
    state["document_names"] = [d["display_name"] for d in documents]
    log(f"  documents ready: {[d['display_name'] for d in documents]}")
    save_state(state)
    log("=== Browser setup complete ===")


def phase_browser_cleanup(api: Api, state: dict[str, Any]) -> None:
    log("=== Browser cleanup ===")
    for key in ("kb_id", "disabled_kb_id"):
        kb_id = state.get(key)
        if not kb_id:
            continue
        response = api.request("GET", f"/api/v1/knowledge-bases/{kb_id}")
        if response.status_code != 200:
            log(f"  {key} already gone")
            continue
        submission = api.request("DELETE", f"/api/v1/knowledge-bases/{kb_id}").json()
        wait_task(api, submission["task_id"], timeout=120)
        log(f"  deleted {kb_id}")
    save_state({})
    STATE_FILE.unlink(missing_ok=True)
    log("=== Browser cleanup complete ===")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("phase", choices=["1", "2", "3", "4", "5", "6", "browser-setup", "browser-cleanup"])
    args = parser.parse_args()
    state = load_state()
    with httpx.Client(timeout=60) as client:
        api = Api(client)
        {
            "1": phase1,
            "2": phase2,
            "3": phase3,
            "4": phase4,
            "5": phase5,
            "6": phase6,
            "browser-setup": phase_browser_setup,
            "browser-cleanup": phase_browser_cleanup,
        }[args.phase](api, state)
    evidence_path = Path(f"/tmp/e2e_profile_evidence_phase{args.phase}.md")
    evidence_path.write_text("\n".join(EVIDENCE) + "\n", encoding="utf-8")
    print(f"\nevidence written to {evidence_path}")


if __name__ == "__main__":
    main()
